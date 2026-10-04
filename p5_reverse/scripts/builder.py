"""Row builder for A02669 reverse new-build batches (schema = Batch120 71 cols + extensions)."""
import openpyxl, json

COLS = ["문항ID","시험ID","원본파일 Drive ID","원본파일명","원본페이지","원문항번호","학교/시험명","학교명","출처대분류","주관기관",
"시행연도","학년도","시행월","학년","교육과정","과목","선택과목","대단원","중단원","소단원","세부유형ID","세부유형명","개념태그","결합유형",
"난이도","난이도근거","배점","문항형식","정답","정답유형","그림/그래프","계산량","조건해석","추론강도","함정포인트","유사문항그룹ID",
"중복대표여부","원문사용가능","변형가능","변형주의","해설파일 Drive ID","해설페이지","해설연결상태","문항이미지/좌표","텍스트추출상태",
"원문시각확인","검산상태","검산근거","최근사용일","사용횟수","학교/프로젝트사용이력","DB신뢰도","메모","분석큐ID","원본SHA256","정답률",
"정답률출처","필드확인필요","분석일","검토상태","정답표 Drive ID","정답표페이지","원본URL","원문요약","상위원문항번호","소문항번호",
"확정정답","원문조건충돌","분석스트림","소문항수","production_ready",
# extension fields (appended; existing order untouched)
"원답지정답","schema_version"]

CFG = dict(SRC_ID="14h8FPJc-DVnRom0HlbVB1ZDvDfoGahRl",
 ZIP_SHA="5b7e36ab5cc427ae3c2c213996f85709a85c2117ca3179e293bf233030bc34d1",
 PDF_SHA="f181427658e8e1605b853215b95e805f507defe6d2f0e192746e9139e394c400",
 TAG="CJPOLY", QID="A02669", NPAGES=22, ANS_PAGES="14~15",
 FNAME="01 다항식(천재-류) 84제 (1).zip > 01 다항식(천재-류) 84제.pdf",
 UNIT_TITLE="1. 다항식", BIG="Ⅰ. 다항식",
 SOL={16:(1,16),17:(17,30),18:(31,43),19:(44,56),20:(57,73),21:(74,83),22:(84,84)}, QUICK=None)
SRC_ID = CFG["SRC_ID"]; ZIP_SHA = CFG["ZIP_SHA"]; PDF_SHA = CFG["PDF_SHA"]; P = ZIP_SHA[:12]
SOL_PAGE = {}
def configure(**kw):
    global SRC_ID, ZIP_SHA, PDF_SHA, P
    CFG.update(kw); SRC_ID = CFG["SRC_ID"]; ZIP_SHA = CFG["ZIP_SHA"]; PDF_SHA = CFG["PDF_SHA"]; P = ZIP_SHA[:12]
    SOL_PAGE.clear()
    for pg, (lo, hi) in CFG["SOL"].items():
        for q in range(lo, hi+1): SOL_PAGE[q] = pg
configure()
def quick_page(q):
    if not CFG["QUICK"]: return CFG["ANS_PAGES"]
    for pg,(lo,hi) in CFG["QUICK"].items():
        if lo<=q<=hi: return str(pg)
    return CFG["ANS_PAGES"]
SECTIONS = {  # (exam suffix, 중단원)
 "S11": ("1-1 다항식의 연산", "다항식의 연산"),
 "S11C": ("1-1 스스로 확인하기", "다항식의 연산"),
 "S12": ("1-2 항등식과 나머지정리", "항등식과 나머지정리"),
 "S12C": ("1-2 스스로 확인하기", "항등식과 나머지정리"),
 "S13": ("1-3 인수분해", "인수분해"),
 "S13C": ("1-3 스스로 확인하기", "인수분해"),
 "MIX": ("스스로 마무리하기(대단원)", "대단원 종합평가"),
 "ACT": ("대단원 활동(모눈 색칠)", "대단원 종합활동"),
}

def row(q, batch, sec, page, small, tid, tname, tags, combo, diff, diff_why, fmt, ans, ans_type,
        calc, cond, infer, trap, summary, check, nsub=1, key=None, fig="없음", visual="직접 확인(원본 130dpi 렌더 판독)",
        status="검산완료", trust="높음(원문 시각판독+정답지+독립풀이 일치)", ready="YES", conflict="", memo_extra="",
        review="분석완료", final=None):
    exam_sfx, mid = SECTIONS[sec]
    sp = SOL_PAGE[q]
    r = {c: "" for c in COLS}
    r.update({
     "문항ID": f"D_{P}_{CFG['TAG']}_Q{q:03d}",
     "시험ID": f"TEXT_{P}_{CFG['TAG']}_{sec}",
     "원본파일 Drive ID": SRC_ID,
     "원본파일명": CFG["FNAME"],
     "원본페이지": str(page), "원문항번호": str(q),
     "학교/시험명": f"천재(류) 수학(상) 교과서 편집본 {CFG['UNIT_TITLE']} — {exam_sfx}",
     "학교명": "해당없음", "출처대분류": "교과서/교과서편집본(학원편집)",
     "주관기관": "천재교육(류희찬) 교과서 기반 · 용문아 수학방 편집",
     "시행연도": "해당없음", "학년도": "해당없음", "시행월": "해당없음",
     "학년": "고1", "교육과정": "2015 개정", "과목": "수학(상)", "선택과목": "해당없음",
     "대단원": CFG["BIG"], "중단원": mid, "소단원": small,
     "세부유형ID": tid, "세부유형명": tname, "개념태그": tags, "결합유형": combo,
     "난이도": diff, "난이도근거": diff_why + " 공식 정답률 미확보.",
     "배점": "미표기", "문항형식": fmt, "정답": ans, "정답유형": ans_type, "그림/그래프": fig,
     "계산량": calc, "조건해석": cond, "추론강도": infer, "함정포인트": trap,
     "유사문항그룹ID": tid, "중복대표여부": "상위 원번호 1개=문항 1행",
     "원문사용가능": "내부참고(원본권리조건 확인)", "변형가능": "가능",
     "변형주의": "계수·차수·조건 변경 시 전체 풀이를 독립 재검산",
     "해설파일 Drive ID": SRC_ID, "해설페이지": f"{quick_page(q)}(빠른정답); {sp}(풀이)",
     "해설연결상태": "동일 PDF 빠른정답+풀이 연결",
     "문항이미지/좌표": f"PDF {page}쪽",
     "텍스트추출상태": "PyMuPDF 본문만 추출(Hancom 수식객체 텍스트 소실) → 원본 렌더 시각판독으로 수식 확정",
     "원문시각확인": visual, "검산상태": status, "검산근거": check,
     "사용횟수": "0", "DB신뢰도": trust,
     "메모": f"{CFG['QID']} {batch}: zip 내부 단일 PDF(SHA256 {PDF_SHA[:16]}…) {CFG['NPAGES']}쪽; " + memo_extra,
     "분석큐ID": CFG["QID"], "원본SHA256": ZIP_SHA, "정답률출처": "미확보",
     "분석일": "2026-10-04", "검토상태": review,
     "정답표 Drive ID": SRC_ID, "정답표페이지": quick_page(q),
     "원본URL": f"https://drive.google.com/file/d/{SRC_ID}/view",
     "원문요약": summary, "상위원문항번호": str(q), "소문항번호": "",
     "확정정답": final if final is not None else ans, "원문조건충돌": conflict,
     "분석스트림": "P5→P4_REVERSE_NEW_BUILD", "소문항수": str(nsub), "production_ready": ready,
     "원답지정답": key if key is not None else ans, "schema_version": "1.2",
    })
    return r

def save(rows, path, summary_pairs):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "요약"
    ws.append(["항목", "값"])
    for k, v in summary_pairs: ws.append([k, v])
    w2 = wb.create_sheet("문항DB"); w2.append(COLS)
    for r in rows:
        w2.append([r[c] for c in COLS])
        for cell in w2[w2.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.data_type = 's'
    wb.save(path)

def validate(rows):
    ids = [r["문항ID"] for r in rows]
    keys = [(r["원본파일 Drive ID"], r["원문항번호"], r["소문항번호"]) for r in rows]
    assert len(ids) == len(set(ids)), "dup id"
    assert len(keys) == len(set(keys)), "dup key"
    bad = [r["문항ID"] for r in rows if r["production_ready"] == "YES" and (not r["세부유형ID"] or not r["확정정답"] or r["원문조건충돌"])]
    assert not bad, bad
    generic = [r["문항ID"] for r in rows if any(g in r["세부유형명"] for g in ["단원 종합", "기타", "미분류"])]
    assert not generic, generic
    return {"rows": len(rows), "subparts": sum(int(r["소문항수"]) for r in rows),
            "YES": sum(r["production_ready"] == "YES" for r in rows),
            "REVIEW/HOLD": sum(r["production_ready"] != "YES" for r in rows)}

def save_csv(rows, path):
    import csv
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(COLS)
        for r in rows: w.writerow([r[c] for c in COLS])
